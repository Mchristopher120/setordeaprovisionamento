"""Importação em lote por planilha Excel (empenhos, itens, pedidos e NFs).

A planilha tem 4 abas (a ordem das colunas é livre; vale o nome do cabeçalho):
  Empenhos: NE | Fornecedor | Processo
  Itens:    NE | Item | Descrição | Unidade | Quantidade | Valor unitário
  Pedidos:  NE | Pedido | Data do pedido | Prazo (dias) | Observação | Item | Qtde pedida
  NFs:      NE | Pedido | NF | Data da NF | Item | Qtde recebida

Pedidos e NFs têm uma linha por item. Linhas com o mesmo NE + Pedido formam
um pedido só; linhas com o mesmo NE + Pedido + NF formam uma NF só.
Pagamento não é importado: a NF entra "a pagar" e é marcada no sistema.
"""
import io
import re
import unicodedata
from datetime import date, datetime

EPS = 1e-9

ABAS = {
    "Empenhos": ["NE", "Fornecedor", "Processo"],
    "Itens": ["NE", "Item", "Descrição", "Unidade", "Quantidade", "Valor unitário"],
    "Pedidos": ["NE", "Pedido", "Data do pedido", "Prazo (dias)", "Observação", "Item", "Qtde pedida"],
    "NFs": ["NE", "Pedido", "NF", "Data da NF", "Item", "Qtde recebida"],
}
OBRIGATORIAS = {
    "Empenhos": ["NE", "Fornecedor"],
    "Itens": ["NE", "Item", "Descrição", "Quantidade", "Valor unitário"],
    "Pedidos": ["NE", "Pedido", "Data do pedido", "Item", "Qtde pedida"],
    "NFs": ["NE", "Pedido", "NF", "Data da NF", "Item", "Qtde recebida"],
}

INSTRUCOES = [
    "COMO PREENCHER",
    "",
    "1. Empenhos: uma linha por empenho (NE, fornecedor e processo).",
    "2. Itens: uma linha por item do empenho. A coluna NE liga o item ao empenho.",
    "3. Pedidos: uma linha por item pedido. Linhas com o mesmo NE + Pedido formam um único pedido.",
    "   A data, o prazo e a observação são lidos da primeira linha do pedido. O prazo pode ficar em branco.",
    "4. NFs: uma linha por item entregue. Linhas com o mesmo NE + Pedido + NF formam uma única NF.",
    "   Pedido ainda não entregue: não coloque linha nesta aba.",
    "",
    "Datas: 18/09/2026 (ou célula de data do Excel). Valores: 1.234,56 ou 1234.56.",
    "O pagamento não é importado: as NFs entram como 'a pagar' e são marcadas como pagas no sistema.",
    "Empenho com NE já cadastrado no sistema não é importado de novo (aparece como erro na verificação).",
    "Nada é gravado enquanto houver erro: o sistema mostra uma prévia e só importa quando você confirmar.",
]


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def normalizar_ne(v):
    return re.sub(r"\s+", "", _texto(v)).upper()


def _data(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = _texto(v)
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s[:10] if fmt == "%Y-%m-%d" else s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _num(v, parse):
    if v is None or _texto(v) == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _texto(v)
    if not re.search(r"\d", s):
        return None
    return parse(s)


def _linhas_da_aba(wb, nome, erros):
    """Lê a aba como lista de (nº da linha, {coluna: valor}). Cabeçalho na 1ª linha não vazia."""
    ws = next((wb[t] for t in wb.sheetnames if _norm(t) == _norm(nome)), None)
    if ws is None:
        if nome in ("Empenhos", "Itens"):
            erros.append(f"Aba “{nome}” não encontrada na planilha.")
        return []
    linhas = list(ws.iter_rows(values_only=True))
    cab_idx = next((i for i, ln in enumerate(linhas) if any(_texto(c) for c in ln)), None)
    if cab_idx is None:
        return []
    cab = {_norm(c): j for j, c in enumerate(linhas[cab_idx]) if _texto(c)}
    mapa = {}
    for col in ABAS[nome]:
        j = cab.get(_norm(col))
        if j is None and col == "Descrição":
            j = cab.get("descricao")
        if j is not None:
            mapa[col] = j
    faltam = [c for c in OBRIGATORIAS[nome] if c not in mapa]
    if faltam:
        erros.append(f"Aba “{nome}”: faltam as colunas {', '.join(faltam)}.")
        return []
    out = []
    for i, ln in enumerate(linhas[cab_idx + 1:], start=cab_idx + 2):
        reg = {c: (ln[j] if j < len(ln) else None) for c, j in mapa.items()}
        if not any(_texto(v) for v in reg.values()):
            continue
        out.append((i, reg))
    return out


def ler_planilha(conteudo, parse_num):
    """Lê e valida a planilha. Devolve (empenhos, erros, avisos).

    Não consulta o banco: a checagem de NE já cadastrado fica com quem chama.
    """
    from openpyxl import load_workbook

    erros, avisos = [], []
    try:
        wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    except Exception:
        return [], ["Não consegui abrir o arquivo. Envie uma planilha Excel (.xlsx)."], []

    def err(aba, linha, msg):
        erros.append(f"{aba}, linha {linha}: {msg}")

    empenhos, por_ne = [], {}
    for ln, r in _linhas_da_aba(wb, "Empenhos", erros):
        ne, forn = normalizar_ne(r.get("NE")), _texto(r.get("Fornecedor")).upper()
        if not ne:
            err("Empenhos", ln, "NE em branco"); continue
        if not forn:
            err("Empenhos", ln, f"fornecedor em branco ({ne})"); continue
        if ne in por_ne:
            err("Empenhos", ln, f"NE {ne} repetido na planilha"); continue
        e = {"ne": ne, "fornecedor": forn, "processo": _texto(r.get("Processo")),
             "itens": [], "pedidos": [], "_itens": {}, "_pedidos": {}}
        por_ne[ne] = e
        empenhos.append(e)

    def empenho_da_linha(aba, ln, r):
        ne = normalizar_ne(r.get("NE"))
        e = por_ne.get(ne)
        if not e:
            err(aba, ln, f"NE {ne or '(em branco)'} não está na aba Empenhos")
        return e

    for ln, r in _linhas_da_aba(wb, "Itens", erros):
        e = empenho_da_linha("Itens", ln, r)
        if not e:
            continue
        num, desc = _texto(r.get("Item")), _texto(r.get("Descrição"))
        qtde, vu = _num(r.get("Quantidade"), parse_num), _num(r.get("Valor unitário"), parse_num)
        if not num:
            err("Itens", ln, "número do item em branco"); continue
        if num in e["_itens"]:
            err("Itens", ln, f"item {num} repetido no empenho {e['ne']}"); continue
        if not desc:
            err("Itens", ln, f"item {num} sem descrição"); continue
        if qtde is None or qtde <= 0:
            err("Itens", ln, f"item {num}: quantidade inválida"); continue
        if vu is None or vu < 0:
            err("Itens", ln, f"item {num}: valor unitário inválido"); continue
        it = {"numero": num, "descricao": desc, "und": _texto(r.get("Unidade")).upper(),
              "qtde": qtde, "vu": vu, "_pedida": 0.0}
        e["_itens"][num] = it
        e["itens"].append(it)

    for ln, r in _linhas_da_aba(wb, "Pedidos", erros):
        e = empenho_da_linha("Pedidos", ln, r)
        if not e:
            continue
        num_ped, num_item = _texto(r.get("Pedido")), _texto(r.get("Item"))
        qtde = _num(r.get("Qtde pedida"), parse_num)
        if not num_ped:
            err("Pedidos", ln, "número do pedido em branco"); continue
        p = e["_pedidos"].get(num_ped)
        if p is None:
            data = _data(r.get("Data do pedido"))
            if not data:
                err("Pedidos", ln, f"pedido {num_ped}: data inválida (use 18/09/2026)"); continue
            prazo_txt = _texto(r.get("Prazo (dias)"))
            prazo = None
            if prazo_txt:
                n = _num(prazo_txt, parse_num)
                if n is None or n < 0 or not float(n).is_integer():
                    err("Pedidos", ln, f"pedido {num_ped}: prazo deve ser um número inteiro de dias")
                else:
                    prazo = int(n)
            p = {"numero": num_ped, "data": data, "prazo": prazo, "obs": _texto(r.get("Observação")),
                 "itens": {}, "nfs": [], "_nfs": {}}
            e["_pedidos"][num_ped] = p
            e["pedidos"].append(p)
        it = e["_itens"].get(num_item)
        if not it:
            err("Pedidos", ln, f"item {num_item or '(em branco)'} não existe no empenho {e['ne']}"); continue
        if qtde is None or qtde <= 0:
            err("Pedidos", ln, f"pedido {num_ped}, item {num_item}: quantidade inválida"); continue
        if num_item in p["itens"]:
            err("Pedidos", ln, f"item {num_item} repetido no pedido {num_ped}"); continue
        p["itens"][num_item] = qtde
        it["_pedida"] += qtde

    for e in empenhos:
        for it in e["itens"]:
            if it["_pedida"] > it["qtde"] + EPS:
                erros.append(f"Empenho {e['ne']}, item {it['numero']}: pedidos somam {it['_pedida']:g}, "
                             f"mais que o empenhado ({it['qtde']:g}).")

    for ln, r in _linhas_da_aba(wb, "NFs", erros):
        e = empenho_da_linha("NFs", ln, r)
        if not e:
            continue
        num_ped, nf, num_item = _texto(r.get("Pedido")), _texto(r.get("NF")), _texto(r.get("Item"))
        qtde = _num(r.get("Qtde recebida"), parse_num)
        p = e["_pedidos"].get(num_ped)
        if not p:
            err("NFs", ln, f"pedido {num_ped or '(em branco)'} não está na aba Pedidos do empenho {e['ne']}"); continue
        if not nf:
            err("NFs", ln, "número da NF em branco"); continue
        n = p["_nfs"].get(nf)
        if n is None:
            data = _data(r.get("Data da NF"))
            if not data:
                err("NFs", ln, f"NF {nf}: data inválida (use 18/09/2026)"); continue
            n = {"nf": nf, "data": data, "itens": {}}
            p["_nfs"][nf] = n
            p["nfs"].append(n)
        if num_item not in p["itens"]:
            err("NFs", ln, f"NF {nf}: item {num_item or '(em branco)'} não faz parte do pedido {num_ped}"); continue
        if qtde is None or qtde <= 0:
            err("NFs", ln, f"NF {nf}, item {num_item}: quantidade inválida"); continue
        n["itens"][num_item] = n["itens"].get(num_item, 0.0) + qtde

    for e in empenhos:
        for p in e["pedidos"]:
            for num_item, pedida in p["itens"].items():
                rec = sum(n["itens"].get(num_item, 0.0) for n in p["nfs"])
                if rec > pedida + EPS:
                    erros.append(f"Empenho {e['ne']}, pedido {p['numero']}, item {num_item}: "
                                 f"NFs somam {rec:g}, mais que o pedido ({pedida:g}).")
        if not e["itens"]:
            avisos.append(f"Empenho {e['ne']} está sem itens.")

    for e in empenhos:
        e.pop("_itens"); e.pop("_pedidos")
        for it in e["itens"]:
            it.pop("_pedida")
        for p in e["pedidos"]:
            p.pop("_nfs")
    return empenhos, erros, avisos


def montar_planilha(empenhos=()):
    """Gera a planilha no formato de importação (modelo vazio ou preenchida)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Instruções"
    for t in INSTRUCOES:
        ws.append([t])
    ws["A1"].font = Font(bold=True, size=13, color="1A3373")
    ws.column_dimensions["A"].width = 110

    linhas = {"Empenhos": [], "Itens": [], "Pedidos": [], "NFs": []}
    for e in empenhos:
        linhas["Empenhos"].append([e["ne"], e["fornecedor"], e.get("processo", "")])
        for it in e["itens"]:
            linhas["Itens"].append([e["ne"], it["numero"], it["descricao"], it.get("und", ""), it["qtde"], it["vu"]])
        for p in e.get("pedidos", []):
            dp = datetime.strptime(p["data"], "%Y-%m-%d")
            for num, q in p["itens"].items():
                linhas["Pedidos"].append([e["ne"], p["numero"], dp, p.get("prazo"), p.get("obs") or None, num, q])
            for n in p.get("nfs", []):
                dn = datetime.strptime(n["data"], "%Y-%m-%d")
                for num, q in n["itens"].items():
                    linhas["NFs"].append([e["ne"], p["numero"], n["nf"], dn, num, q])

    formatos = {
        "Itens": {5: "#,##0.###", 6: '"R$" #,##0.00'},
        "Pedidos": {3: "DD/MM/YYYY", 7: "#,##0.###"},
        "NFs": {4: "DD/MM/YYYY", 6: "#,##0.###"},
    }
    larguras = {"Descrição": 70, "Observação": 40, "Fornecedor": 34}
    for nome, cab in ABAS.items():
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1A3373")
            c.alignment = Alignment(horizontal="center", vertical="center")
        for ln in linhas[nome]:
            ws.append(ln)
        for col, fmt in formatos.get(nome, {}).items():
            for cell in ws[get_column_letter(col)][1:]:
                cell.number_format = fmt
        for i, h in enumerate(cab, 1):
            ws.column_dimensions[get_column_letter(i)].width = larguras.get(h, max(12, len(h) + 4))
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
